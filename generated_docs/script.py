from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import os, textwrap, re, math

outdir="generated_data"
report_docx=os.path.join(outdir,"MERA_RF_Event_Recorder_종합보고서.docx")
spec_docx=os.path.join(outdir,"MERA_RF_Event_Recorder_HAS_스펙문서.docx")
report_pdf=os.path.join(outdir,"MERA_RF_Event_Recorder_종합보고서.pdf")
spec_pdf=os.path.join(outdir,"MERA_RF_Event_Recorder_HAS_스펙문서.pdf")

# ---------- shared content ----------
title="MERA — Multi-channel Deterministic RF Event Recorder IP"
subtitle="AMD Versal RF 고객 RTL 공백 분석 및 독립 IP 하드웨어 아키텍처 명세서"
date="2026-09-22"

sources = [
("AMD Versal RF Series Product Page", "https://www.amd.com/en/products/adaptive-socs-and-fpgas/versal/rf-series.html"),
("AMD DS950 — Versal Architecture and Product Data Sheet", "https://docs.amd.com/api/khub/documents/BY0FMkLH05y85Nwu4GmWsw/content"),
("AMD DS945 — Versal XQ Architecture and Product Data Sheet", "https://docs.amd.com/api/khub/documents/wBgkyKdTzQhGgubawOKewA/content"),
("AMD AM009 — AI Engine Architecture Manual: Stream Switch Buffering", "https://docs.amd.com/r/en-US/am009-versal-ai-engine/Stream-Switch-Buffering-and-Latency"),
("AMD AM009 — AI Engine to PL Interface", "https://docs.amd.com/r/en-US/am009-versal-ai-engine/AI-Engine-to-Programmable-Logic-Interface"),
("AMD AM009 — AXI4-Stream Interconnect", "https://docs.amd.com/r/en-US/am009-versal-ai-engine/AXI4-Stream-Interconnect"),
("AMD PG022 — AXI DataMover", "https://docs.amd.com/r/en-US/pg022_axi_datamover/Features"),
("AMD/Xilinx RFSoC-MTS GitHub", "https://github.com/Xilinx/RFSoC-MTS"),
("AMD/Xilinx RFDC interrupt example", "https://github.com/Xilinx/embeddedsw/blob/master/XilinxProcessorIPLib/drivers/rfdc/examples/xrfdc_intr_example.c"),
]

def set_cell_shading(cell, fill="D9EAF7"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), "true")
    trPr.append(tblHeader)

def style_doc(doc):
    styles=doc.styles
    styles["Normal"].font.name="Arial"
    styles["Normal"].font.size=Pt(10)
    for s in ["Title","Heading 1","Heading 2","Heading 3"]:
        styles[s].font.name="Arial"
    styles["Title"].font.size=Pt(22)
    styles["Heading 1"].font.size=Pt(16)
    styles["Heading 2"].font.size=Pt(13)
    styles["Heading 3"].font.size=Pt(11)

def add_table(doc, headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers))
    t.style="Table Grid"
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    hdr=t.rows[0]
    set_repeat_table_header(hdr)
    for i,h in enumerate(headers):
        hdr.cells[i].text=str(h)
        set_cell_shading(hdr.cells[i])
        hdr.cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row):
            cells[i].text=str(v)
            cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths):
                row.cells[i].width=Inches(w)
    return t

def add_bullets(doc, items, level=0):
    for item in items:
        p=doc.add_paragraph(style="List Bullet" if level==0 else "List Bullet 2")
        p.add_run(item)

def add_numbered(doc, items):
    for item in items:
        p=doc.add_paragraph(style="List Number")
        p.add_run(item)

# ---------- REPORT DOCX ----------
doc=Document()
style_doc(doc)
sec=doc.sections[0]
sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.7); sec.right_margin=Inches(.7)

p=doc.add_paragraph()
p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run(title); r.bold=True; r.font.size=Pt(24)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.add_run(subtitle).italic=True
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.add_run(f"문서일: {date}\n상태: Concept / Architecture Baseline\n대상 플랫폼: AMD Versal RF 계열, 특히 VR1952")

doc.add_page_break()

doc.add_heading("Executive Summary",1)
doc.add_paragraph(
"본 보고서는 AMD Versal RF 계열에서 AMD가 이미 제공하는 RF 데이터 변환, DSP hard IP, AI Engine/NoC/AXI 데이터 이동 기능을 먼저 분리한 뒤, "
"그 위에서 고객이 반복적으로 작성해야 하는 application-specific RTL의 공백을 찾는 것을 목적으로 한다. "
"결론은 범용 Data Mover, FFT, Channelizer, Resampler 또는 단순 FIFO/Backpressure IP를 독립 제품으로 만드는 것은 차별화 근거가 약하다는 것이다."
)
doc.add_paragraph(
"대안으로 제안하는 MERA(Multi-channel Deterministic RF Event Recorder)는 여러 RF 채널의 스트림에서 이벤트를 기준으로 "
"pre-trigger/post-trigger 구간을 결정적으로 보존하고, 채널 간 sample index를 정렬하며, timestamp/metadata를 포함한 event record를 "
"AXI4-MM을 통해 DDR에 burst write하는 application-layer IP이다."
)
doc.add_paragraph(
"중요한 한계도 명확히 한다. 공개 자료만으로 'AMD가 동일한 capture IP를 제품으로 절대 제공하지 않는다'고 증명할 수는 없다. "
"그러나 AMD/Xilinx 공개 RFSoC 설계에서는 capture memory, trigger, FIFO flush, deep capture, DMA가 별도의 설계 특화 블록으로 나타나며, "
"RFDC 드라이버 예제도 capture/stimulus register write가 design-specific이라고 명시한다. 따라서 본 보고서는 '공개 자료에서 확인되는 고객 RTL 영역'을 "
"제품 후보로 좁힌 것이다."
)

doc.add_heading("1. 문제 정의와 조사 질문",1)
add_numbered(doc,[
"AMD가 VR1952에서 이미 silicon/hard IP로 해결하는 기능은 무엇인가?",
"그 기능 사이의 경계에서 고객이 실제로 RTL을 작성해야 하는 부분은 무엇인가?",
"그 고객 RTL이 단순 glue logic인지, 반복 가능한 독립 IP로 묶을 가치가 있는지?",
"FPGA/HLS 환경에서도 검증 가능하고, 향후 design-house/IP 업무 포트폴리오로 연결할 수 있는 구조인가?"
])

doc.add_heading("2. 기준 제품: AMD Versal RF VR1952",1)
add_table(doc,
["항목","VR1952 공개 사양","근거"],
[
["14-bit RF ADC","8개, 최대 32 GSPS","AMD 제품 페이지 / DS950"],
["14-bit RF DAC","16개, 최대 16 GSPS","AMD 제품 페이지 / DS950"],
["AI Engine","120 tiles","AMD 제품 페이지 / DS950"],
["LDPC Decoder","6","AMD 제품 페이지 / DS950"],
["Channelizer","공식 제품 페이지 320개; DS950 리소스 표는 96 × 1 GSPS로 표기","문서 버전/카운트 정의 차이. DS950를 자원량 기준으로 사용"],
["FFT/iFFT","공식 제품 페이지 40개; DS950 리소스 표는 36개","문서 버전/카운트 정의 차이. DS950를 자원량 기준으로 사용"],
["Poly Block","24","AMD 제품 페이지 / DS950"],
["System Logic Cells","2,473,800","DS950"],
["LUT","1,130,880","DS950"],
["DSP58","3,976","DS950"],
["PL Memory","189 Mb","AMD 제품 페이지 / DS950"],
["DDR Memory Controllers","5","AMD 제품 페이지 / DS950"],
["GTM/GTMP","20 GTMP, 56G(112G) 표기","AMD 제품 페이지 / DS950"],
],
[1.6,3.2,2.0])

doc.add_paragraph(
"주의: AMD의 공개 웹 제품 페이지와 최신 DS950에서 Channelizer/FFT 개수가 다르게 보인다. "
"이는 문서 버전 또는 '논리 블록 수'와 '특정 속도/구성 기준 자원 수'의 정의 차이 가능성이 있으므로 본 설계에서는 DS950를 "
"정확한 현재 리소스 기준으로 삼고, 웹 페이지 수치는 별도 참고치로만 기록한다."
)

doc.add_heading("3. AMD가 이미 해결하는 영역",1)
add_table(doc,
["영역","확인된 기능/수치","판정"],
[
["RF data conversion","14-bit ADC/DAC, DDC/DUC 등","AMD 제공"],
["FFT/iFFT","4 GSPS; 8~4096 point; 18-bit input / 31-bit output","독립 FFT IP 기회 낮음"],
["Channelizer","64-tap programmable filter + 8-point FFT","독립 channelizer 기회 낮음"],
["Polyphase arbitrary resampler","256 phases, 17 taps/phase, 4,352-tap prototype; 1 GSPS/instance","독립 resampler 기회 낮음"],
["AIE stream switch","32-bit AXI4-Stream crossbar; FIFO/arbiter/backpressure","범용 stream infrastructure 중복"],
["PL→AIE","8 streams/column; 64-bit physical stream, 32/64-bit configurable; aggregate 32 GB/s/column","범용 boundary IP 중복"],
["AIE→PL","6 streams/column; aggregate 24 GB/s/column","범용 boundary IP 중복"],
["AIE buffering","PL→AIE CDC FIFO 12-deep; switch FIFO 4-deep/port; optional 16-deep chainable FIFO","범용 FIFO IP 중복"],
["AXI DataMover","8~1024-bit AXI4-Stream; 32~1024-bit AXI4; burst 2~256 beats; MM2S/S2MM","범용 DMA IP 중복"],
],
[1.55,4.2,1.25])

doc.add_heading("4. 데이터 경계별 Gap Analysis",1)
add_table(doc,
["Boundary","Width/Clock","Throughput","Buffer/Backpressure","고객 RTL","AMD 해결"],
[
["RF ADC → PL","VR1952 exact public interface width는 본 조사에서 미확정","ADC rate/decimation 의존","RFDC 내부/인터페이스 buffering 존재 가능; exact VR1952 공개 수치 미확정","application routing/selection/capture","RF ADC + DDC/DUC"],
["PL → FFT","18b in / 31b out; IP 4 GSPS","4 GSPS/IP","IP 내부 상세 공개 범위 제한","stream connection / framing","hard FFT/iFFT"],
["PL → Channelizer","세부 physical interface 공개값 제한","1 GSPS class resource","IP 내부 상세 제한","stream connection / channel mapping","hard channelizer"],
["PL → Resampler","16b complex / 24b coeff","1 GSPS/instance; 4 GSPS configurations","IP 내부 상세 제한","stream connection / control","hard polyphase resampler"],
["PL ↔ AIE","64b physical; 32/64 configurable; two streams→128b 구성 가능","PL→AIE 32 GB/s/column; AIE→PL 24 GB/s/column","12-deep CDC + 4-deep switch + optional 16-deep","application dataflow","AIE interconnect/FIFO"],
["Stream ↔ DDR","DataMover 8~1024b stream, 32~1024b MM","AXI/memory dependent","store-forward/queues/burst","command generation / record policy","DataMover"],
["Event capture","애플리케이션 정의","event rate dependent","customer-defined circular/pre/post buffers","trigger policy, event record, alignment, timestamp","공개 자료상 generic primitive만 확인"],
],
[1.25,2.0,1.45,2.0,2.2,1.6])

doc.add_heading("5. 결정적 근거: 공개 RFSoC 설계에서 실제로 남는 RTL",1)
doc.add_paragraph(
"AMD/Xilinx의 공개 RFSoC-MTS 설계는 3개 ADC 채널을 내부 memory에 64 kilosamples 저장하고, 네 번째 채널을 deep-capture 모듈로 보내며, "
"더 긴 capture에는 PL-DRAM을 사용할 수 있다고 설명한다."
)
add_bullets(doc,[
"ADC capture memories가 채널별로 존재한다.",
"trig_cap 신호가 ADC capture를 실제로 trigger한다.",
"fifo_flush가 DMA FIFO를 초기화한다.",
"adc_dma가 PL DMA를 통해 DDR4로 데이터를 이동한다.",
"deepCapture 모듈이 장시간 capture 경로로 사용된다."
])
doc.add_paragraph(
"더 직접적인 근거로 AMD/Xilinx RFDC interrupt example은 stimulus/capture register write에 대해 "
"'Below writes are not generic, they are design specific'이라고 명시하고, capture block register가 user-specific design에 따라 달라질 수 있음을 설명한다."
)

doc.add_heading("6. 왜 단순 Capture IP가 아닌가",1)
doc.add_paragraph(
"단순 ADC→BRAM→DDR capture는 상품 차별성이 약하다. MERA의 제품 경계는 '데이터 저장'이 아니라 "
"'이벤트를 중심으로 여러 채널의 시간 구간을 결정적으로 구성하는 것'이다."
)
add_bullets(doc,[
"Pre-trigger circular buffer: 이벤트 발생 이전 데이터를 보존",
"Post-trigger capture: 이벤트 이후 지정된 길이까지 수집",
"Multi-channel sample-index alignment: 여러 채널을 동일 기준점으로 묶음",
"Timestamp: event 및 record 생성 시점의 시간 정보",
"Record formatter: header + channel metadata + sample payload",
"Deterministic burst writer: DDR에 정해진 형식으로 기록",
"Overflow/discontinuity detection: 데이터 유실 여부를 명시적으로 기록",
"Event queue: 여러 이벤트가 연속 발생해도 record를 잃지 않도록 정책화"
])

doc.add_heading("7. 제안 IP의 시스템 위치",1)
doc.add_paragraph(
"권장 시스템 구성은 다음과 같다."
)
doc.add_paragraph(
"RF ADC → DDC/Channelizer/FFT/AIE detector → MERA → AXI4-MM/DataMover → DDR"
)
doc.add_paragraph(
"MERA는 AMD hard IP를 대체하지 않는다. 오히려 AMD hard IP에서 생성된 '관심 이벤트'와 고속 sample stream을 "
"제품 수준의 기록 단위로 바꾸는 상위 application IP다."
)

doc.add_heading("8. 제안 아키텍처",1)
add_table(doc,
["블록","역할","핵심 상태/구현"],
[
["AXIS Input Adapter","입력 stream 표준화","TVALID/TREADY/TLAST/TKEEP 처리"],
["Channel Demux","채널별 데이터 분리","channel ID/stream mapping"],
["Circular Pre-Trigger Buffer","최근 N samples 유지","BRAM/URAM 기반 ring buffer"],
["Trigger Engine","capture 시작 조건 결정","external/software/threshold/event"],
["Capture FSM","pre→trigger→post→commit 상태 전이","IDLE/PRE/ARMED/POST/COMMIT/ERROR"],
["Timestamp Counter","64-bit timebase","free-running counter"],
["Alignment Engine","채널 sample index 정렬","common sample counter + offset table"],
["Record Formatter","메타데이터 + payload 구성","fixed header + channel descriptors"],
["Event Queue","복수 이벤트 관리","descriptor FIFO"],
["AXI-MM Writer","DDR burst write","burst scheduler / address manager"],
["Status/Telemetry","overflow, drop, latency 등","AXI4-Lite registers"],
])

doc.add_heading("9. 권장 v1.0 하드웨어 사양",1)
add_table(doc,
["Parameter","권장값","허용/비고"],
[
["Channels","8","4/8/16 parameterizable"],
["Input sample","Complex I/Q","v1.0 기본"],
["I/Q width","16-bit + 16-bit","parameterizable 8~24b"],
["AXI-S input width","128-bit","32/64/128 parameterizable"],
["AXI-S clock","250 MHz baseline","실제 FPGA 결과에 따라 200~500 MHz sweep"],
["Timestamp","64-bit","free-running"],
["Pre-trigger depth","64 Ki samples/channel","2^10~2^20 권장 parameter"],
["Post-trigger depth","256 Ki samples/channel","2^10~2^22 parameter"],
["Trigger sources","4","EXT / SW / threshold / upstream event"],
["Event ID","32-bit","monotonic"],
["Channel alignment","sample-index exact","±0 sample 목표"],
["Output","AXI4-MM S2MM","AMD DataMover와 연결 가능"],
["Record header","64 bytes","versioned"],
["Max events in queue","64","parameterizable"],
["Overflow policy","drop-new + sticky flag","v1.0 deterministic"],
["CRC","optional 32-bit","record-level"],
])

doc.add_heading("10. Record Format v1.0",1)
add_table(doc,
["Offset","Bytes","Field","설명"],
[
["0x00","4","MAGIC","0x4D455241 ('MERA')"],
["0x04","2","VERSION","0x0100"],
["0x06","2","HEADER_LEN","64"],
["0x08","4","EVENT_ID","event sequence"],
["0x0C","4","CHANNEL_MASK","활성 채널"],
["0x10","8","TIMESTAMP","trigger sample time"],
["0x18","4","PRE_SAMPLES","trigger 이전 sample 수"],
["0x1C","4","POST_SAMPLES","trigger 이후 sample 수"],
["0x20","4","SAMPLE_WIDTH","I/Q bit width"],
["0x24","4","SAMPLE_FORMAT","complex/fixed-point encoding"],
["0x28","4","ALIGN_STATUS","alignment result"],
["0x2C","4","STATUS","overflow/drop/error bits"],
["0x30","8","PAYLOAD_BYTES","payload size"],
["0x38","8","RESERVED","future use"],
],
[.75,.65,1.35,4.2])

doc.add_heading("11. Register Map v1.0",1)
add_table(doc,
["Offset","Register","R/W","내용"],
[
["0x000","CONTROL","RW","ENABLE/ARM/SOFT_TRIGGER/FLUSH"],
["0x004","STATUS","RO","ARMED/TRIGGERED/CAPTURING/DONE/OVERFLOW/ERROR"],
["0x008","TRIGGER_SEL","RW","trigger source select"],
["0x00C","CHANNEL_MASK","RW","active channels"],
["0x010","PRE_LEN","RW","pre-trigger samples"],
["0x014","POST_LEN","RW","post-trigger samples"],
["0x018","TIMESTAMP_LO","RO","timestamp low"],
["0x01C","TIMESTAMP_HI","RO","timestamp high"],
["0x020","EVENT_ID","RO","current event id"],
["0x024","WRITE_ADDR_LO","RW","DDR base low"],
["0x028","WRITE_ADDR_HI","RW","DDR base high"],
["0x02C","MAX_RECORDS","RW","event queue limit"],
["0x030","IRQ_ENABLE","RW","interrupt mask"],
["0x034","IRQ_STATUS","RW1C","interrupt status"],
["0x038","ALIGN_OFFSET0","RW","CH0 alignment offset"],
["0x03C","ALIGN_OFFSET1","RW","CH1 alignment offset"],
["0x040","ALIGN_OFFSET2","RW","CH2 alignment offset"],
["0x044","ALIGN_OFFSET3","RW","CH3 alignment offset"],
["0x048","ALIGN_OFFSET4","RW","CH4 alignment offset"],
["0x04C","ALIGN_OFFSET5","RW","CH5 alignment offset"],
["0x050","ALIGN_OFFSET6","RW","CH6 alignment offset"],
["0x054","ALIGN_OFFSET7","RW","CH7 alignment offset"],
["0x058","CAP_COUNT","RO","completed capture count"],
["0x05C","DROP_COUNT","RO","dropped event count"],
])

doc.add_heading("12. FSM 명세",1)
add_table(doc,
["State","진입 조건","동작","종료 조건"],
[
["RESET","reset","register init, pointer clear","reset deassert"],
["IDLE","not armed","stream ignored/optional monitor","ARM"],
["ARMED","ARM=1","circular pre-buffer 계속 갱신","TRIGGER"],
["TRIGGERED","trigger accepted","trigger timestamp/event id latch","next cycle"],
["POST_CAPTURE","triggered","post-trigger samples 수집","post_len reached"],
["COMMIT","capture complete","record header/payload descriptor 생성","descriptor queued"],
["DMA_WAIT","descriptor queued","AXI-MM writer arbitration","write complete"],
["DONE","write complete","status/IRQ set","next event/REARM"],
["ERROR","overflow/invalid config","sticky error latch","software clear"],
])

doc.add_heading("13. Backpressure 및 결정성 정책",1)
doc.add_paragraph(
"입력 AXI4-Stream의 TREADY가 내려갈 수 있는 구조는 이벤트 기록에서 데이터 유실 위험을 만든다. "
"따라서 v1.0의 권장 정책은 정상 capture 구간에서는 입력을 backpressure하지 않고, 내부 circular buffer가 항상 sample rate를 흡수하도록 한다."
)
add_bullets(doc,[
"Normal armed mode: TREADY=1을 목표로 유지",
"DDR writer가 느려져도 capture buffer가 우선 데이터를 보존",
"buffer watermark가 임계치를 넘으면 STATUS에 EARLY_OVERFLOW 경고",
"실제 overflow 발생 시 DROP/OVERFLOW sticky bit 기록",
"event commit과 DDR write는 분리하여 capture path의 실시간성을 보장",
"DMA/DDR contention은 record queue가 흡수"
])

doc.add_heading("14. 성능 모델",1)
doc.add_paragraph(
"8채널, I/Q 각 16-bit라면 채널당 32 bit/sample이다. 8채널 aggregate raw payload는 "
"8 × 32 = 256 bit/sample = 32 byte/sample이다."
)
add_table(doc,
["샘플률","8채널 raw payload","필요 지속 write BW"],
[
["100 MSPS","3.2 GB/s","최소 3.2 GB/s + overhead"],
["250 MSPS","8.0 GB/s","최소 8.0 GB/s + overhead"],
["500 MSPS","16.0 GB/s","최소 16.0 GB/s + overhead"],
["1 GSPS","32.0 GB/s","최소 32.0 GB/s + overhead"],
])
doc.add_paragraph(
"이 표는 MERA의 외부 DDR이 실제로 해당 bandwidth를 보장한다는 뜻이 아니다. "
"이는 설계 요구량을 산정하기 위한 raw payload 모델이다. 실제 sustained BW는 DDR controller, NoC, AXI width, burst length, arbitration, "
"clock, 다른 master의 traffic에 의해 결정되므로 FPGA 보드에서 측정해야 한다."
)

doc.add_heading("15. 검증 계획",1)
add_table(doc,
["검증 레벨","시험","Pass 기준"],
[
["C-sim","trigger 위치별 pre/post payload","golden reference와 100% 동일"],
["C/RTL co-sim","record header/payload","bit-exact"],
["RTL simulation","AXI backpressure","no illegal transfer / no lost sample under specified buffer margin"],
["Formal/Assertion","FSM safety","invalid state/illegal transition 없음"],
["FPGA","100/250/500 MHz sweep","timing closure 및 measured throughput"],
["FPGA","4/8/16 channel sweep","resource scaling 기록"],
["FPGA","DDR contention","defined traffic 조건에서 drop policy 검증"],
["Fault injection","trigger storm","DROP_COUNT/STATUS 정확"],
["Alignment test","채널별 artificial offset","보정 후 ±0 sample"],
])

doc.add_heading("16. PPA/Engineering Ablation",1)
add_bullets(doc,[
"Architecture A: BRAM circular buffer + single AXI writer",
"Architecture B: per-channel buffer + shared formatter",
"Architecture C: packed interleaved buffer + shared writer",
"Architecture D: event descriptor queue depth 16/32/64/128",
"AXI width: 64/128/256/512 bit",
"burst length: 16/32/64/128/256 beats",
"pre-trigger depth: 4K/16K/64K/256K",
"channel count: 4/8/16"
])
doc.add_paragraph(
"논문의 핵심은 '새로운 ML 알고리즘'이 아니라, 동일한 capture requirement를 만족시키는 hardware architecture 사이에서 "
"resource, Fmax, latency, sustained write bandwidth, determinism의 trade-off를 계측하는 것이다."
)

doc.add_heading("17. 상품화 경계와 비목표",1)
add_table(doc,
["포함","비포함"],
[
["event trigger/record policy","ADC analog frontend"],
["multi-channel alignment","FFT/channelizer 자체"],
["pre/post circular capture","generic AXI DataMover 자체"],
["timestamp/metadata","DDR controller 자체"],
["record formatting","AIE compute kernel"],
["capture overflow diagnostics","특정 radar detection algorithm"],
["AXI4-S/AXI4-MM integration","특정 RF application의 proprietary detector"],
])

doc.add_heading("18. 시장/제품 주장에 대한 증거 수준",1)
add_table(doc,
["주장","증거 수준","판정"],
[
["AMD는 FFT/channelizer/resampler/DMA를 이미 제공","높음","공식 문서 확인"],
["AMD AIE interconnect가 FIFO/backpressure를 제공","높음","AM009 확인"],
["공개 RFSoC 설계에서 capture/trigger/deep capture가 별도 구성","높음","Xilinx 공개 GitHub 확인"],
["RFDC example의 capture register가 design-specific","높음","AMD/Xilinx 코드 주석 직접 확인"],
["모든 Versal RF 고객이 이 문제를 반복적으로 겪는다","낮음~중간","공개 자료만으로 정량 증명 불가"],
["AMD가 MERA와 동일한 제품을 판매하지 않는다","미확정","비공개/early-access IP 가능성 때문에 단정 금지"],
["독립 상용 IP로 충분한 시장가치가 있다","검증 필요","고객 인터뷰/benchmark 필요"],
])

doc.add_heading("19. 개발 로드맵",1)
add_numbered(doc,[
"Stage 0 — AMD 공개 reference design에서 capture path의 RTL/driver boundary를 추가 조사",
"Stage 1 — 4-channel, 128-bit AXIS, BRAM circular buffer, software trigger",
"Stage 2 — pre/post capture + timestamp + record header",
"Stage 3 — 8-channel alignment + external/threshold trigger",
"Stage 4 — AXI-MM burst writer + DDR FPGA measurement",
"Stage 5 — HLS vs hand RTL architecture/resource comparison",
"Stage 6 — parameter sweep와 reproducible benchmark package",
"Stage 7 — 고객 요구를 가정한 product datasheet/API/verification collateral 작성"
])

doc.add_heading("20. 최종 결론",1)
doc.add_paragraph(
"현재까지 조사된 근거를 종합하면, 독립 IP의 초점을 AMD의 primitive를 대체하는 방향으로 잡는 것은 설득력이 약하다. "
"특히 FFT, channelizer, arbitrary resampler, AIE stream switch, FIFO/backpressure, AXI DataMover는 이미 상당히 강력하다."
)
doc.add_paragraph(
"반면 공개 설계에서 반복적으로 나타나는 capture/trigger/deep-capture/data-recording 영역은 application-specific logic으로 남아 있다. "
"따라서 MERA는 'AMD IP를 더 빠르게 복제하는 IP'가 아니라 'AMD hard IP가 생산한 RF stream을 deterministic event record로 제품화하는 IP'로 정의하는 것이 타당하다."
)
doc.add_paragraph(
"단, 상용화 전제의 가장 중요한 다음 검증은 AMD의 비공개 customer IP 존재 여부가 아니라, 공개/실제 고객 설계에서 동일 기능이 반복적으로 수백~수천 줄의 RTL/driver glue로 재구현되는지, "
"그리고 MERA가 이를 얼마나 줄이는지 정량화하는 것이다."
)

doc.add_heading("참고 문헌 / 공식 자료",1)
for name,url in sources:
    p=doc.add_paragraph(style="List Bullet")
    p.add_run(name+" — ").bold=True
    p.add_run(url)

doc.save(report_docx)

# ---------- SPEC DOCX ----------
spec=Document()
style_doc(spec)
sec=spec.sections[0]
sec.top_margin=Inches(.55); sec.bottom_margin=Inches(.55); sec.left_margin=Inches(.6); sec.right_margin=Inches(.6)

p=spec.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run("HAS-MERA-260922"); r.bold=True; r.font.size=Pt(14)
p=spec.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run("MERA v1.0 Hardware Architecture Specification"); r.bold=True; r.font.size=Pt(23)
p=spec.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Multi-channel Deterministic RF Event Recorder IP\nRevision 1.0 — 2026-09-22")

spec.add_page_break()

spec.add_heading("0. Document Control",1)
add_table(spec,["항목","값"],[
["Document ID","HAS-MERA-260922"],
["IP Name","MERA — Multi-channel Deterministic RF Event Recorder"],
["Version","v1.0"],
["Status","Architecture Baseline / RTL Entry Candidate"],
["Target","AMD Versal RF class / FPGA prototype"],
["Implementation","HLS C++ + RTL wrapper/critical control as required"],
["Primary interfaces","AXI4-Stream input, AXI4-Lite control, AXI4-MM output"],
["Primary objective","deterministic multi-channel event capture and record formation"],
["Sign-off condition","all MUST requirements in Section 4 pass"],
])

spec.add_heading("1. Scope",1)
spec.add_paragraph(
"MERA는 이미 존재하는 RF ADC/DSP/AIE primitive를 대체하지 않는다. "
"입력 stream을 이벤트 중심의 multi-channel record로 변환하고, pre-trigger/post-trigger 데이터를 보존하며, "
"timestamp와 상태를 부가하고, DDR로 지속 저장할 수 있는 capture/control IP를 정의한다."
)
spec.add_heading("2. Non-Goals",1)
add_bullets(spec,[
"RF ADC analog control/Calibration 구현",
"FFT/channelizer/resampler 구현",
"AI Engine kernel 구현",
"generic AXI DataMover 재구현",
"DDR controller 구현",
"특정 radar/EMSO detection algorithm 구현",
"아날로그 trigger comparator 구현"
])

spec.add_heading("3. External Architecture",1)
spec.add_paragraph("권장 top-level:")
spec.add_paragraph(
"RF/DSP/AIE stream → AXIS Input Adapter → Channel/Alignment → Circular Capture → Trigger/FSM → Record Formatter → Event Queue → AXI-MM Writer → DDR"
)
spec.add_paragraph("Control path: AXI4-Lite → CSR block → Trigger/FSM/Writer/Status")

spec.add_heading("4. Normative Requirements",1)
reqs=[
("MERA-REQ-001","MUST","AXI4-Stream compliant input handshaking"),
("MERA-REQ-002","MUST","Armed 상태에서 정의된 정상 입력률을 buffer가 흡수해야 한다."),
("MERA-REQ-003","MUST","Trigger timestamp를 64-bit로 latch해야 한다."),
("MERA-REQ-004","MUST","Pre-trigger와 post-trigger sample 수를 독립적으로 설정할 수 있어야 한다."),
("MERA-REQ-005","MUST","활성 채널의 sample index alignment 결과를 record에 기록해야 한다."),
("MERA-REQ-006","MUST","완료된 record는 versioned header를 가져야 한다."),
("MERA-REQ-007","MUST","DDR write failure/overflow를 sticky status로 표시해야 한다."),
("MERA-REQ-008","MUST","Event ID가 동일 event를 유일하게 식별해야 한다."),
("MERA-REQ-009","MUST","AXI4-Lite CSR로 ARM/DISARM/TRIGGER/FLUSH가 가능해야 한다."),
("MERA-REQ-010","MUST","Configuration error를 capture 시작 전에 감지해야 한다."),
("MERA-REQ-011","SHOULD","4/8/16 channel parameterization"),
("MERA-REQ-012","SHOULD","128/256/512-bit AXIS packing 옵션"),
("MERA-REQ-013","SHOULD","External trigger와 software trigger 동시 지원"),
("MERA-REQ-014","SHOULD","threshold/event input 지원"),
("MERA-REQ-015","SHOULD","record-level CRC32"),
("MERA-REQ-016","MAY","dual-buffer DMA scheduling"),
]
add_table(spec,["ID","Priority","Requirement"],reqs)

spec.add_heading("5. Parameter Table",1)
add_table(spec,["Parameter","Default","Range/Constraint"],[
["N_CH","8","4/8/16"],
["IQ_BITS","16","8~24"],
["AXIS_W","128","32/64/128/256/512"],
["TS_W","64","fixed"],
["PRE_DEPTH","65536","power-of-two"],
["POST_DEPTH","262144","power-of-two"],
["EVENT_Q_DEPTH","64","16~256"],
["TRIG_SOURCES","4",">=2"],
["HEADER_BYTES","64","fixed v1.0"],
["ADDR_W","64","32/64"],
["MAX_BURST_BEATS","64","2/4/8/16/32/64/128/256; backend dependent"],
])

spec.add_heading("6. Interface Specification",1)
spec.add_heading("6.1 AXI4-Stream Input",2)
add_table(spec,["Signal","Dir","Width","Requirement"],[
["ACLK","in","1","input clock"],
["ARESETN","in","1","active-low reset"],
["TDATA","in","AXIS_W","sample payload"],
["TVALID","in","1","source valid"],
["TREADY","out","1","sink ready"],
["TLAST","in","1","optional frame boundary"],
["TKEEP","in","AXIS_W/8","byte validity"],
["TUSER","in",">=1","optional channel/event metadata"],
])
spec.add_paragraph("v1.0은 TDATA packing convention을 명시적으로 고정한다: [channel][I][Q] 또는 interleaved channel-major. 실제 RTL sign-off 전에 하나를 선택하고 golden model과 일치시킨다.")

spec.add_heading("6.2 AXI4-Lite Control",2)
spec.add_paragraph("32-bit register data, 32-bit aligned addressing을 기본으로 한다.")

spec.add_heading("6.3 AXI4-MM Output",2)
add_table(spec,["항목","v1.0"],[
["Direction","M_AXI_S2MM"],
["Address","64-bit"],
["Data width","128-bit baseline"],
["Burst","up to 64 beats baseline; backend allows 2~256"],
["Outstanding","parameterized"],
["Alignment","record base aligned to AXI width"],
["Write policy","record contiguous, header first"],
])

spec.add_heading("7. Internal Block Specifications",1)
for h,txt in [
("7.1 AXIS Input Adapter","AXIS handshake를 검증하고 sample beat를 내부 canonical format으로 변환한다. TKEEP/TLAST illegal pattern은 ERROR로 기록한다."),
("7.2 Channel Demux","TUSER 또는 static packing rule에 따라 채널별 sample index를 생성한다."),
("7.3 Sample Counter","capture clock 기준 free-running 64-bit sample counter. Trigger acceptance cycle의 counter 값을 timestamp로 latch한다."),
("7.4 Circular Buffer","각 채널의 최근 PRE_DEPTH samples를 유지한다. write pointer는 modulo power-of-two wrap을 사용한다."),
("7.5 Trigger Engine","EXT, SW, threshold/event input을 OR/priority policy로 조합한다. 이미 POST_CAPTURE 중인 경우 새로운 trigger 처리 정책은 DROP/QUEUE로 설정 가능해야 한다."),
("7.6 Alignment Engine","채널별 offset register와 common sample counter를 이용해 common event index를 생성한다. 허용 범위를 벗어나면 ALIGN_ERROR를 설정한다."),
("7.7 Capture FSM","RESET→IDLE→ARMED→TRIGGERED→POST_CAPTURE→COMMIT→DMA_WAIT→DONE. Error는 별도 sticky state."),
("7.8 Record Formatter","64-byte header와 payload를 결합한다. header version을 통해 향후 포맷 확장을 허용한다."),
("7.9 Event Queue","capture completion descriptor를 FIFO에 넣어 DDR writer와 capture path를 분리한다."),
("7.10 AXI-MM Writer","descriptor를 읽어 header/payload를 burst write한다. DDR contention이 있어도 capture path가 즉시 정지하지 않도록 queueing한다."),
("7.11 Status/Telemetry","overflow, drop, alignment, invalid config, DMA error, event count를 sticky/live 형태로 제공한다.")
]:
    spec.add_heading(h,2); spec.add_paragraph(txt)

spec.add_heading("8. Capture State Machine",1)
add_table(spec,["State","Output/Action","Transition"],[
["RESET","clear state/pointers/status","ARESETN=1"],
["IDLE","no capture","ARM=1→ARMED"],
["ARMED","circular buffer active","trigger→TRIGGERED"],
["TRIGGERED","latch timestamp/event id; freeze logical pre-window","next→POST_CAPTURE"],
["POST_CAPTURE","collect POST_LEN","count==POST_LEN→COMMIT"],
["COMMIT","write descriptor/header metadata","queue available→DMA_WAIT"],
["DMA_WAIT","DDR write","complete→DONE; error→ERROR"],
["DONE","IRQ/status update","REARM→ARMED / else IDLE"],
["ERROR","sticky error","SW CLEAR→IDLE"],
])

spec.add_heading("9. Timing and Determinism",1)
spec.add_paragraph(
"Trigger acceptance latency와 event timestamp의 정의를 분리한다. timestamp는 trigger가 acceptance된 clock cycle의 sample counter 값으로 정의한다. "
"따라서 trigger source 자체의 asynchronous propagation delay는 외부 시스템 측정 항목이며, IP 내부의 acceptance-to-latch latency는 cycle-count로 고정한다."
)
add_table(spec,["Metric","Definition","Target"],[
["T_accept","trigger valid → trigger accepted","1~N cycles, source별 고정"],
["T_latch","accepted → timestamp latch","1 cycle"],
["T_post","trigger → post capture complete","POST_LEN samples + fixed control latency"],
["T_commit","capture complete → descriptor enqueue","bounded constant"],
["T_dma","descriptor → DDR completion","memory/interconnect dependent"],
["Alignment error","max |CH_i index - common index|","0 sample after configured compensation"],
])

spec.add_heading("10. Buffer Sizing Equations",1)
spec.add_paragraph("복소 I/Q interleaved sample의 channel당 byte/sample은 2×IQ_BITS/8. 총 pre-buffer 용량:")
spec.add_paragraph("B_pre = N_CH × PRE_DEPTH × (2×IQ_BITS/8) bytes")
spec.add_paragraph("post-buffer 용량:")
spec.add_paragraph("B_post = N_CH × POST_DEPTH × (2×IQ_BITS/8) bytes")
spec.add_paragraph("예: N_CH=8, IQ_BITS=16, PRE=65536이면 B_pre = 8×65536×4 = 2,097,152 bytes = 2 MiB.")
spec.add_paragraph("POST=262144이면 B_post = 8×262144×4 = 8,388,608 bytes = 8 MiB.")

spec.add_heading("11. Record Memory Layout",1)
add_table(spec,["Field","Size","Encoding"],[
["MAGIC","4B","ASCII MERA"],
["VERSION","2B","0x0100"],
["HEADER_LEN","2B","64"],
["EVENT_ID","4B","unsigned"],
["CHANNEL_MASK","4B","bit per channel"],
["TIMESTAMP","8B","unsigned sample time"],
["PRE_SAMPLES","4B","unsigned"],
["POST_SAMPLES","4B","unsigned"],
["SAMPLE_WIDTH","4B","bits"],
["SAMPLE_FORMAT","4B","enum"],
["ALIGN_STATUS","4B","bitfield"],
["STATUS","4B","bitfield"],
["PAYLOAD_BYTES","8B","unsigned"],
["RESERVED","8B","zero"],
["PAYLOAD","variable","channel-major or agreed packing"],
])

spec.add_heading("12. CSR Register Specification",1)
add_table(spec,["Addr","Name","Access","Bit/Field","Reset","Description"],[
["0x000","CONTROL","RW","0 ENABLE","0","enable block"],
["0x000","CONTROL","RW","1 ARM","0","arm capture"],
["0x000","CONTROL","RW","2 SW_TRIGGER","0","1-cycle pulse"],
["0x000","CONTROL","RW","3 FLUSH","0","clear buffers/queue"],
["0x004","STATUS","RO","0 ARMED","0","armed"],
["0x004","STATUS","RO","1 TRIGGERED","0","trigger latched"],
["0x004","STATUS","RO","2 CAPTURING","0","post capture"],
["0x004","STATUS","RO","3 DONE","0","record done"],
["0x004","STATUS","RO","4 OVERFLOW","0","sticky"],
["0x004","STATUS","RO","5 ALIGN_ERROR","0","sticky"],
["0x004","STATUS","RO","6 DMA_ERROR","0","sticky"],
["0x008","TRIGGER_SEL","RW","[3:0]","0","source select"],
["0x00C","CHANNEL_MASK","RW","N_CH bits","0","active channels"],
["0x010","PRE_LEN","RW","31:0","0","pre samples"],
["0x014","POST_LEN","RW","31:0","0","post samples"],
["0x018","TS_LO","RO","31:0","0","timestamp low"],
["0x01C","TS_HI","RO","31:0","0","timestamp high"],
["0x020","EVENT_ID","RO","31:0","0","current event"],
["0x024","BASE_LO","RW","31:0","0","DDR base low"],
["0x028","BASE_HI","RW","31:0","0","DDR base high"],
["0x02C","EVENT_Q_DEPTH","RW","15:0","64","queue depth"],
["0x030","IRQ_EN","RW","31:0","0","interrupt enables"],
["0x034","IRQ_STATUS","RW1C","31:0","0","interrupt status"],
["0x038-0x054","ALIGN_OFFSET[i]","RW","signed","0","per-channel sample offset"],
["0x058","CAP_COUNT","RO","31:0","0","completed captures"],
["0x05C","DROP_COUNT","RO","31:0","0","dropped events"],
])

spec.add_heading("13. Error Model",1)
add_table(spec,["Code","Condition","Action"],[
["E001","invalid PRE/POST configuration","reject ARM"],
["E002","channel mask empty","reject ARM"],
["E003","buffer overflow","sticky + record status"],
["E004","event queue full","DROP_COUNT++"],
["E005","alignment mismatch","ALIGN_ERROR"],
["E006","AXI write response error","DMA_ERROR"],
["E007","illegal AXIS TKEEP/TLAST","STREAM_ERROR"],
["E008","timestamp discontinuity","TIMESTAMP_ERROR"],
])

spec.add_heading("14. HLS Partitioning Recommendation",1)
add_table(spec,["Block","HLS 적합성","이유"],[
["Record Formatter","높음","stream transform + fixed header"],
["Channel Packing","높음","array/stream processing"],
["Trigger FSM","중간","simple FSM, HLS 가능"],
["Circular Buffer","중간","memory banking/partitioning 제어 필요"],
["AXI-MM Writer","중간","burst semantics와 outstanding 제어가 중요"],
["CSR","낮음~중간","얇은 RTL wrapper 권장"],
["Assertions/Protocol","RTL","SVA 및 protocol checker가 적합"],
])
spec.add_paragraph(
"권장 구현은 '100% HLS'가 아니라 HLS datapath + RTL shell이다. 특히 AXI4-Lite CSR, 정확한 AXI protocol boundary, "
"reset/interrupt, SVA는 RTL로 고정하고, record packing/formatting/copy path는 HLS로 구현하면 설계 의도가 명확하다."
)

spec.add_heading("15. Verification Plan",1)
add_table(spec,["Test ID","Scenario","Expected"],[
["V001","no trigger for 1M samples","no record"],
["V002","SW trigger at sample 1000","pre/post exact"],
["V003","trigger at circular wrap","correct wrap reconstruction"],
["V004","trigger immediately after ARM","short/zero pre policy exact"],
["V005","backpressure burst","no illegal AXIS behavior"],
["V006","event storm","queue/drop policy exact"],
["V007","4/8/16 channel","channel map exact"],
["V008","per-channel offset ±1/±4","alignment correction exact"],
["V009","DDR write error","DMA_ERROR sticky"],
["V010","reset during capture","no ghost record after reset"],
["V011","max PRE/POST","address range safe"],
["V012","random trigger positions","10^5 randomized events, bit-exact golden comparison"],
])

spec.add_heading("16. Acceptance Criteria",1)
add_bullets(spec,[
"100,000 random trigger simulation에서 golden model과 payload mismatch = 0",
"AXI4-Stream protocol assertion failure = 0",
"invalid configuration이 ARM을 차단",
"4/8/16 channel parameter build 성공",
"FPGA target에서 지정 clock constraint에 timing closure",
"pre-trigger wrap boundary에서 sample loss = 0 under specified buffer margin",
"alignment test에서 compensation 후 0-sample error",
"DMA error 및 event drop가 software-visible status로 정확히 보고됨"
])

spec.add_heading("17. FPGA Benchmark Matrix",1)
add_table(spec,["Case","Channels","AXIS","Clock","PRE","POST","측정"],[
["B1","4","128","250MHz","4K","16K","LUT/BRAM/Fmax/latency"],
["B2","8","128","250MHz","64K","256K","LUT/BRAM/Fmax/latency"],
["B3","16","256","250MHz","64K","256K","LUT/BRAM/Fmax/latency"],
["B4","8","256","500MHz","64K","256K","sustained BW"],
["B5","8","512","250MHz","64K","256K","DDR efficiency"],
["B6","8","128","250MHz","256K","1M","memory scaling"],
])

spec.add_heading("18. Known Limitations / Open Questions",1)
add_bullets(spec,[
"VR1952의 exact RF-ADC→PL physical interface width/clock/buffering 공개 수치는 본 HAS에서 확정하지 않는다. 제품별 confidential documentation으로 sign-off 필요.",
"AMD 제품 페이지와 DS950 사이의 Channelizer/FFT count discrepancy는 문서 정의/버전 차이로 보이며, RTL 설계 자원 산정은 DS950 기준으로 한다.",
"MERA와 기능적으로 동일한 AMD customer/early-access IP 존재 여부는 공개 자료로 확정할 수 없다.",
"DDR sustained bandwidth는 특정 FPGA board와 memory configuration에서 실측해야 한다.",
"threshold trigger를 MERA 내부에서 계산할지 AIE/PL upstream detector가 제공할지는 제품 profile에 따라 분리할 수 있다.",
])

spec.add_heading("19. Sign-off Checklist",1)
add_table(spec,["항목","Owner","Status"],[
["Architecture reviewed","Lead","OPEN"],
["Interface packing frozen","RTL/HLS","OPEN"],
["CSR map frozen","RTL","OPEN"],
["Record format frozen","SW/HLS","OPEN"],
["Golden model available","Verification","OPEN"],
["C-sim passed","Verification","OPEN"],
["Co-sim passed","Verification","OPEN"],
["FPGA timing passed","Implementation","OPEN"],
["DDR throughput measured","System","OPEN"],
["Customer-equivalent benchmark completed","Architecture","OPEN"],
])

spec.add_heading("20. References",1)
for name,url in sources:
    p=spec.add_paragraph(style="List Bullet")
    p.add_run(name+" — ").bold=True
    p.add_run(url)

spec.save(spec_docx)

# ---------- PDFs ----------
# Try Korean font
font_paths=[
    "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
    "/usr/share/fonts/opentype/noto/NotoSansCJKkr-Regular.otf",
    "/usr/share/fonts/truetype/nanum/NanumGothic.ttf"
]
font_path=next((p for p in font_paths if os.path.exists(p)),None)
if font_path:
    pdfmetrics.registerFont(TTFont("Korean",font_path))
    base_font="Korean"
else:
    base_font="Helvetica"

styles=getSampleStyleSheet()
styles.add(ParagraphStyle(name="KTitle", parent=styles["Title"], fontName=base_font, fontSize=20, leading=25, alignment=1, spaceAfter=10))
styles.add(ParagraphStyle(name="KH1", parent=styles["Heading1"], fontName=base_font, fontSize=14, leading=18, spaceBefore=12, spaceAfter=7))
styles.add(ParagraphStyle(name="KH2", parent=styles["Heading2"], fontName=base_font, fontSize=11.5, leading=15, spaceBefore=9, spaceAfter=5))
styles.add(ParagraphStyle(name="KBody", parent=styles["BodyText"], fontName=base_font, fontSize=8.6, leading=12, spaceAfter=5))
styles.add(ParagraphStyle(name="KSmall", parent=styles["BodyText"], fontName=base_font, fontSize=7.3, leading=9.5))
styles.add(ParagraphStyle(name="KTable", parent=styles["BodyText"], fontName=base_font, fontSize=6.8, leading=8.5))

def esc(s):
    return str(s).replace("&","&amp;").replace("<","&lt;").replace(">","&gt;")

def build_pdf(path, title_text, sections):
    docpdf=SimpleDocTemplate(path,pagesize=A4,rightMargin=28,leftMargin=28,topMargin=28,bottomMargin=28)
    story=[Paragraph(esc(title_text),styles["KTitle"]), Paragraph(f"{date}",styles["KBody"]), Spacer(1,8)]
    for kind,data in sections:
        if kind=="h1": story.append(Paragraph(esc(data),styles["KH1"]))
        elif kind=="h2": story.append(Paragraph(esc(data),styles["KH2"]))
        elif kind=="p": story.append(Paragraph(esc(data),styles["KBody"]))
        elif kind=="bullets":
            for x in data: story.append(Paragraph("• "+esc(x),styles["KBody"]))
        elif kind=="table":
            headers,rows=data
            tab=[[Paragraph("<b>"+esc(h)+"</b>",styles["KTable"]) for h in headers]]
            for row in rows:
                tab.append([Paragraph(esc(v),styles["KTable"]) for v in row])
            t=Table(tab, repeatRows=1, hAlign="LEFT")
            t.setStyle(TableStyle([
                ("GRID",(0,0),(-1,-1),0.35,colors.grey),
                ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#D9EAF7")),
                ("VALIGN",(0,0),(-1,-1),"TOP"),
                ("LEFTPADDING",(0,0),(-1,-1),3),
                ("RIGHTPADDING",(0,0),(-1,-1),3),
                ("TOPPADDING",(0,0),(-1,-1),3),
                ("BOTTOMPADDING",(0,0),(-1,-1),3),
            ]))
            story.append(t); story.append(Spacer(1,6))
    docpdf.build(story)

# Build concise-but-detailed PDFs from main docs
report_sections=[
("h1","Executive Summary"),
("p","AMD Versal RF VR1952에서 범용 FFT/channelizer/resampler/AIE interconnect/DataMover를 독립 IP로 대체하는 방향은 차별화 근거가 약하다. 공개 설계에서 반복적으로 나타나는 capture/trigger/deep-capture/data-recording 영역을 상위 application IP인 MERA로 묶는 것이 현재 가장 근거가 좋은 후보이다."),
("h1","VR1952 기준 사양"),
("table",(["항목","VR1952","근거"],[
["RF ADC","8 × 14-bit, 32 GSPS","AMD DS950 / product page"],
["RF DAC","16 × 14-bit, 16 GSPS","AMD DS950 / product page"],
["AIE","120 tiles","AMD DS950 / product page"],
["LDPC","6 decoders","AMD DS950 / product page"],
["Channelizer","DS950: 96 × 1 GSPS; web page: 320","문서 정의/버전 차이"],
["FFT/iFFT","DS950: 36; web page: 40","문서 정의/버전 차이"],
["Poly block","24","AMD DS950"],
["LUT","1,130,880","AMD DS950"],
["DSP58","3,976","AMD DS950"],
])),
("h1","AMD가 이미 제공하는 것"),
("bullets",[
"4 GSPS FFT/iFFT, 8~4096 point, 18-bit input/31-bit output.",
"64-tap + 8-point FFT channelizer.",
"Polyphase arbitrary resampler: 256 phases, 17 taps/phase, 4,352-tap prototype; 1 GSPS/instance.",
"AIE stream switch: 32-bit AXI4-Stream crossbar, FIFO, arbiter, backpressure.",
"PL→AIE 32 GB/s/column, AIE→PL 24 GB/s/column.",
"PL→AIE CDC FIFO 12-deep, stream-switch FIFO 4-deep/port, optional 16-deep FIFO.",
"AXI DataMover: stream 8~1024b, MM 32~1024b, burst 2~256 beats."
]),
("h1","공개 설계에서 확인되는 고객 영역"),
("p","Xilinx RFSoC-MTS 공개 설계는 3개 ADC 채널의 64 kilosample internal capture, 네 번째 채널의 deep capture, PL-DRAM 장기 저장, trigger, FIFO flush, DMA를 함께 사용한다. AMD/Xilinx RFDC interrupt example은 stimulus/capture register write가 design specific이라고 직접 명시한다."),
("h1","MERA 제품 정의"),
("p","Multi-channel Deterministic RF Event Recorder. RF/DSP/AIE stream을 받아 trigger를 기준으로 pre-trigger/post-trigger 구간을 보존하고, channel alignment, 64-bit timestamp, event ID, status를 record header에 넣어 AXI4-MM/DDR로 저장한다."),
("h1","v1.0 핵심 스펙"),
("table",(["항목","기본값"],[
["Channels","8"],
["I/Q","16+16 bit complex"],
["AXI-S","128 bit"],
["Timestamp","64 bit"],
["Pre-trigger","64 Ki samples/channel"],
["Post-trigger","256 Ki samples/channel"],
["Event queue","64"],
["Trigger","external/software/threshold/upstream event"],
["Output","AXI4-MM S2MM"],
["Header","64 bytes"],
["Alignment","0-sample residual target"],
])),
("h1","검증 및 상품화 판단"),
("bullets",[
"4/8/16 channel, 128/256/512-bit AXIS, 4K~1M sample buffer sweep.",
"FPGA에서 LUT/BRAM/URAM/Fmax/latency/sustained DDR BW 측정.",
"100,000 randomized triggers bit-exact golden comparison.",
"상용화 전에는 AMD confidential/early-access IP 존재 여부와 실제 고객 반복 RTL 규모를 별도로 확인해야 한다."
]),
("h1","참고자료"),
("bullets",[x[0]+" — "+x[1] for x in sources])
]
build_pdf(report_pdf,title,report_sections)

spec_sections=[
("h1","문서 목적"),
("p","MERA v1.0은 deterministic multi-channel RF event capture를 위한 architecture baseline이다. AXI4-Stream input, AXI4-Lite control, AXI4-MM output을 기본으로 한다."),
("h1","Normative Requirements"),
("table",(["ID","Priority","Requirement"],[
["REQ-001","MUST","AXI4-Stream handshake 준수"],
["REQ-002","MUST","ARMED 상태에서 지정 입력률 흡수"],
["REQ-003","MUST","64-bit trigger timestamp latch"],
["REQ-004","MUST","독립 pre/post length"],
["REQ-005","MUST","channel sample-index alignment 결과 기록"],
["REQ-006","MUST","versioned record header"],
["REQ-007","MUST","overflow/DMA error sticky status"],
["REQ-008","MUST","unique Event ID"],
["REQ-009","MUST","ARM/DISARM/TRIGGER/FLUSH CSR"],
["REQ-010","MUST","invalid configuration 사전 검출"],
["REQ-011","SHOULD","4/8/16 channel"],
["REQ-012","SHOULD","128/256/512-bit AXIS"],
["REQ-013","SHOULD","external/software trigger"],
])),
("h1","Parameter"),
("table",(["Parameter","Default","Range"],[
["N_CH","8","4/8/16"],["IQ_BITS","16","8~24"],["AXIS_W","128","32~512"],["TS_W","64","fixed"],
["PRE_DEPTH","65536","power-of-two"],["POST_DEPTH","262144","power-of-two"],["EVENT_Q_DEPTH","64","16~256"],
["HEADER_BYTES","64","fixed"],["ADDR_W","64","32/64"]
])),
("h1","Architecture"),
("bullets",[
"AXIS Input Adapter",
"Channel Demux / Alignment",
"Circular Pre-trigger Buffer",
"Trigger Engine",
"Capture FSM",
"64-bit Timestamp Counter",
"Record Formatter",
"Event Descriptor Queue",
"AXI4-MM Burst Writer",
"CSR / Status / Telemetry"
]),
("h1","Record Header"),
("table",(["Offset","Bytes","Field"],[
["0x00","4","MAGIC"],["0x04","2","VERSION"],["0x06","2","HEADER_LEN"],["0x08","4","EVENT_ID"],
["0x0C","4","CHANNEL_MASK"],["0x10","8","TIMESTAMP"],["0x18","4","PRE_SAMPLES"],["0x1C","4","POST_SAMPLES"],
["0x20","4","SAMPLE_WIDTH"],["0x24","4","SAMPLE_FORMAT"],["0x28","4","ALIGN_STATUS"],["0x2C","4","STATUS"],
["0x30","8","PAYLOAD_BYTES"],["0x38","8","RESERVED"]
])),
("h1","FSM"),
("table",(["State","Function"],[
["IDLE","capture off"],["ARMED","circular buffer running"],["TRIGGERED","timestamp/event id latch"],
["POST_CAPTURE","post samples collect"],["COMMIT","descriptor creation"],["DMA_WAIT","DDR burst write"],
["DONE","status/IRQ"],["ERROR","sticky error"]
])),
("h1","Verification"),
("bullets",[
"Random trigger 100,000 cases, bit-exact golden model.",
"AXI protocol assertions.",
"Wrap-boundary tests.",
"4/8/16 channel tests.",
"±1/±4 sample channel offset alignment tests.",
"DDR write error injection.",
"FPGA timing and sustained bandwidth measurement."
]),
("h1","Open Questions"),
("bullets",[
"VR1952 exact RF-ADC→PL physical interface details require product-specific documentation.",
"AMD public product page vs DS950 has Channelizer/FFT count discrepancies; DS950 is used for current resource accounting.",
"Exact absence of an equivalent AMD confidential/customer IP cannot be proven from public sources.",
"DDR sustained bandwidth is board/system dependent."
]),
("h1","References"),
("bullets",[x[0]+" — "+x[1] for x in sources])
]
build_pdf(spec_pdf,"HAS-MERA-260922 — MERA v1.0 Hardware Architecture Specification",spec_sections)

print("생성 완료:")
for p in [report_docx,spec_docx,report_pdf,spec_pdf]:
    print(p, os.path.getsize(p))
